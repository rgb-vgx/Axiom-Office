from __future__ import annotations

import os
import tempfile
from copy import copy as _copy

from docx import Document


def to_json(obj) -> str:
    import json
    return json.dumps(obj, ensure_ascii=False, default=str)


def _atomic_save(doc, path: str) -> None:
    directory = os.path.dirname(os.path.abspath(path))
    handle, temp_path = tempfile.mkstemp(
        prefix=".~" + os.path.basename(path) + ".", suffix=".tmp", dir=directory
    )
    os.close(handle)
    try:
        doc.save(temp_path)
    except Exception:
        try:
            os.remove(temp_path)
        except OSError:
            pass
        raise
    try:
        os.replace(temp_path, path)
    except PermissionError as exc:
        try:
            os.remove(temp_path)
        except OSError:
            pass
        raise PermissionError(
            f"cannot write '{path}': file is locked (open in Word/WPS?) - close it first or use word_type_text on the live bridge"
        ) from exc


def profile(path: str) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    doc = Document(path)
    style_counts: dict = {}
    for paragraph in doc.paragraphs:
        name = paragraph.style.name if paragraph.style is not None else "None"
        style_counts[name] = style_counts.get(name, 0) + 1
    core = doc.core_properties
    return {
        "path": path,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "sections": len(doc.sections),
        "style_counts": style_counts,
        "core": {
            "title": core.title,
            "author": core.author,
            "created": core.created,
            "modified": core.modified,
        },
    }


def get_text(path: str, max_chars: int = 0, include_tables: bool = True) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    doc = Document(path)
    parts = [paragraph.text for paragraph in doc.paragraphs]
    if include_tables:
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(parts)
    total = len(text)
    truncated = False
    if max_chars and len(text) > max_chars:
        text = text[:max_chars]
        truncated = True
    return {"path": path, "chars": total, "truncated": truncated, "text": text}


def find_text(path: str, query: str, max_results: int = 50) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    doc = Document(path)
    needle = (query or "").lower()
    matches: list = []
    for index, paragraph in enumerate(doc.paragraphs):
        if needle and needle in paragraph.text.lower():
            matches.append({
                "index": index,
                "style": paragraph.style.name if paragraph.style is not None else "None",
                "text": paragraph.text[:300],
            })
            if len(matches) >= max_results:
                break
    return {"query": query, "count": len(matches), "matches": matches}


def extract_table(path: str, index: int = 0) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    doc = Document(path)
    tables = doc.tables
    if index < 0 or index >= len(tables):
        raise ValueError(f"table index out of range (file has {len(tables)} tables)")
    table = tables[index]
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    return {"path": path, "index": index, "row_count": len(rows), "rows": rows}


def create(path: str, paragraphs: list | None = None, title: str | None = None,
           author: str | None = None, overwrite: bool = False) -> dict:
    if os.path.exists(path) and not overwrite:
        raise ValueError("file already exists - pass overwrite=true to replace it")
    doc = Document()
    added = 0
    for spec in paragraphs or []:
        if isinstance(spec, str):
            doc.add_paragraph(spec)
            added += 1
            continue
        if not isinstance(spec, dict):
            continue
        if "table" in spec:
            data = spec.get("table") or []
            if data:
                table = doc.add_table(rows=len(data), cols=len(data[0]))
                table.style = "Table Grid"
                for r, row in enumerate(data):
                    for c, value in enumerate(row):
                        table.cell(r, c).text = "" if value is None else str(value)
                added += 1
            continue
        text = spec.get("text", "")
        style = spec.get("style")
        if style:
            doc.add_paragraph(text, style=style)
        else:
            doc.add_paragraph(text)
        added += 1
    if title:
        doc.core_properties.title = title
    if author:
        doc.core_properties.author = author
    _atomic_save(doc, path)
    return {"created": path, "blocks": added}
