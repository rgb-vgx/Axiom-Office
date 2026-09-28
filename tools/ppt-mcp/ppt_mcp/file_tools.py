from __future__ import annotations

import json
import os
import tempfile

from pptx import Presentation
from pptx.util import Inches


def to_json(obj) -> str:
    return json.dumps(obj, ensure_ascii=False, default=str)


def _atomic_save(prs, path: str) -> None:
    directory = os.path.dirname(os.path.abspath(path))
    handle, temp_path = tempfile.mkstemp(
        prefix=".~" + os.path.basename(path) + ".", suffix=".tmp", dir=directory
    )
    os.close(handle)
    try:
        prs.save(temp_path)
    except Exception:
        try:
            os.remove(temp_path)
        except OSError:
            pass
        raise
    try:
        os.replace(temp_path, path)
    except PermissionError as exc:
        try:
            os.remove(temp_path)
        except OSError:
            pass
        raise PermissionError(
            f"cannot write '{path}': file is locked (open in PowerPoint/WPS?) - close it first or use the live bridge"
        ) from exc


def profile(path: str) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    prs = Presentation(path)
    slides = []
    for index, slide in enumerate(prs.slides, start=1):
        texts = []
        shapes = 0
        for shape in slide.shapes:
            shapes += 1
            try:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    texts.append(shape.text_frame.text)
            except Exception:
                continue
        notes = ""
        try:
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text or ""
        except Exception:
            notes = ""
        slides.append({
            "index": index,
            "layout": slide.slide_layout.name,
            "shapes": shapes,
            "texts": texts,
            "notes": notes[:300],
        })
    return {
        "path": path,
        "slide_count": len(prs.slides._sldIdLst),
        "size": {"width": prs.slide_width, "height": prs.slide_height},
        "slides": slides,
    }


def get_text(path: str, max_chars: int = 0) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    prs = Presentation(path)
    parts = []
    for index, slide in enumerate(prs.slides, start=1):
        parts.append(f"[Slide {index}]")
        for shape in slide.shapes:
            try:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    parts.append(shape.text_frame.text)
            except Exception:
                continue
        try:
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("[Notes] " + slide.notes_slide.notes_text_frame.text)
        except Exception:
            continue
    text = "\n".join(parts)
    total = len(text)
    truncated = False
    if max_chars and len(text) > max_chars:
        text = text[:max_chars]
        truncated = True
    return {"path": path, "chars": total, "truncated": truncated, "text": text}


def create(path: str, slides: list | None = None, overwrite: bool = False) -> dict:
    if os.path.exists(path) and not overwrite:
        raise ValueError("file already exists - pass overwrite=true to replace it")
    prs = Presentation()
    created = 0
    for spec in slides or []:
        if isinstance(spec, str):
            spec = {"title": spec}
        if not isinstance(spec, dict):
            continue
        layout_index = int(spec.get("layout", 1))
        layout_index = max(0, min(layout_index, len(prs.slide_layouts) - 1))
        slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
        title_text = spec.get("title")
        if title_text:
            try:
                slide.shapes.title.text = str(title_text)
            except Exception:
                pass
        bullets = spec.get("bullets")
        if bullets is None and spec.get("text") is not None:
            bullets = [str(spec["text"])]
        if bullets:
            body = None
            for placeholder in slide.placeholders:
                try:
                    if placeholder.placeholder_format.idx == 1:
                        body = placeholder
                        break
                except Exception:
                    continue
            if body is not None:
                frame = body.text_frame
                frame.text = str(bullets[0])
                for extra in bullets[1:]:
                    paragraph = frame.add_paragraph()
                    paragraph.text = str(extra)
        created += 1
    _atomic_save(prs, path)
    return {"created": path, "slides": created}


def add_slide(path: str, title: str | None = None, bullets: list | None = None,
              layout: int = 1) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    prs = Presentation(path)
    layout_index = max(0, min(int(layout), len(prs.slide_layouts) - 1))
    slide = prs.slides.add_slide(prs.slide_layouts[layout_index])
    if title:
        try:
            slide.shapes.title.text = str(title)
        except Exception:
            pass
    if bullets:
        for placeholder in slide.placeholders:
            try:
                if placeholder.placeholder_format.idx == 1:
                    frame = placeholder.text_frame
                    frame.text = str(bullets[0])
                    for extra in bullets[1:]:
                        frame.add_paragraph().text = str(extra)
                    break
            except Exception:
                continue
    _atomic_save(prs, path)
    return {"path": path, "slide_count": len(prs.slides._sldIdLst)}
