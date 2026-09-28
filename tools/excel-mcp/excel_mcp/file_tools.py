from __future__ import annotations

import csv
import datetime as _dt
import json
import math
import os
import re
import tempfile
from copy import copy
from typing import Any

import duckdb
import openpyxl
import pandas as pd
import xlrd
from openpyxl.styles import PatternFill, Side
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries
from openpyxl.worksheet.table import Table, TableStyleInfo

OPENPYXL_FORMATS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
XLRD_FORMATS = {".xls"}
CSV_FORMATS = {".csv", ".tsv", ".txt"}


def to_json(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, default=str)


def _ext(path: str) -> str:
    return os.path.splitext(path)[1].lower()


def _check_format(path: str) -> str:
    ext = _ext(path)
    if ext in OPENPYXL_FORMATS or ext in XLRD_FORMATS or ext in CSV_FORMATS:
        return ext
    raise ValueError(f"unsupported format '{ext}': supported are xlsx, xlsm, xls, csv, tsv")


def _detect_kind(path: str) -> str:
    try:
        with open(path, "rb") as handle:
            head = handle.read(8)
    except OSError:
        head = b""
    if head.startswith(b"PK"):
        return "xlsx"
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return "xls"
    ext = _ext(path)
    if ext in OPENPYXL_FORMATS:
        return "xlsx"
    if ext in XLRD_FORMATS:
        return "xls"
    return "csv"


def _openpyxl_load(path: str, read_only: bool, data_only: bool, keep_vba: bool = False):
    ext = _ext(path)
    if ext in OPENPYXL_FORMATS:
        return openpyxl.load_workbook(path, read_only=read_only, data_only=data_only, keep_vba=keep_vba), None
    stream = open(path, "rb")
    try:
        wb = openpyxl.load_workbook(stream, read_only=read_only, data_only=data_only, keep_vba=keep_vba)
        return wb, stream
    except Exception:
        stream.close()
        raise


def _clean_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (_dt.datetime, _dt.date, _dt.time)):
        return value.isoformat()
    return value


def _clean_matrix(rows) -> list[list]:
    return [[_clean_value(v) for v in row] for row in rows]


def _coerce_csv_cell(text: str) -> Any:
    stripped = text.strip()
    if stripped == "":
        return ""
    try:
        as_int = int(stripped)
        if str(as_int) == stripped:
            return as_int
    except ValueError:
        pass
    try:
        as_float = float(stripped)
        if repr(as_float) == stripped or str(as_float) == stripped:
            return as_float
    except ValueError:
        pass
    return text


def _csv_dialect_for(path: str, sample: str):
    if _ext(path) == ".tsv":
        return csv.excel_tab
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return csv.excel


def _read_csv_rows(path: str) -> list[list]:
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        dialect = _csv_dialect_for(path, sample)
        return [[_coerce_csv_cell(cell) for cell in row] for row in csv.reader(handle, dialect)]


def _xls_sheet(book: xlrd.book.Book, sheet: str | None):
    if sheet is None:
        return book.sheet_by_index(0)
    try:
        return book.sheet_by_name(sheet)
    except xlrd.biffh.XLRDError as exc:
        raise ValueError(f"sheet not found: {sheet}") from exc


def _xls_cell_value(book: xlrd.book.Book, cell) -> Any:
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
        except (ValueError, xlrd.XLDateError):
            return cell.value
    if cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
        return None
    return cell.value


def _xls_rows_range(path: str, sheet: str | None, min_row: int, max_row: int,
                    min_col: int, max_col: int, has_header_row: bool = True):
    book = xlrd.open_workbook(path)
    ws = _xls_sheet(book, sheet)
    rows: list[list] = []
    start = min_row if has_header_row else min_row
    for r in range(max(0, start - 1), min(ws.nrows, max_row)):
        row: list = []
        for c in range(max(0, min_col - 1), min(ws.ncols, max_col)):
            row.append(_xls_cell_value(book, ws.cell(r, c)))
        rows.append(row)
    return ws.name, rows, ws.nrows, ws.ncols


def _sheet_names(path: str) -> list[str]:
    ext = _check_format(path)
    if ext in OPENPYXL_FORMATS:
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    if ext in XLRD_FORMATS:
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            return list(book.sheet_names())
        finally:
            book.release_resources()
    return [os.path.splitext(os.path.basename(path))[0]]


def profile(path: str, sheet: str | None = None) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    _check_format(path)
    kind = _detect_kind(path)
    result: dict = {"path": path, "format": kind, "sheets": []}

    if kind == "csv":
        if sheet and sheet != os.path.splitext(os.path.basename(path))[0]:
            raise ValueError("csv files have a single sheet")
        rows = _read_csv_rows(path)
        header = rows[0] if rows else []
        sample = rows[1:6] if rows else []
        result["sheets"].append({
            "name": os.path.splitext(os.path.basename(path))[0],
            "rows": len(rows),
            "columns": max((len(r) for r in rows), default=0),
            "header": [_clean_value(v) for v in header],
            "sample": _clean_matrix(sample),
        })
        return result

    if kind == "xls":
        book = xlrd.open_workbook(path)
        names = [sheet] if sheet else book.sheet_names()
        for name in names:
            ws = _xls_sheet(book, name)
            header: list = []
            sample: list = []
            if ws.nrows > 0 and ws.ncols > 0:
                header = [_clean_value(_xls_cell_value(book, ws.cell(0, c))) for c in range(ws.ncols)]
                for r in range(1, min(ws.nrows, 6)):
                    sample.append([_clean_value(_xls_cell_value(book, ws.cell(r, c))) for c in range(ws.ncols)])
            result["sheets"].append({
                "name": ws.name,
                "rows": ws.nrows,
                "columns": ws.ncols,
                "header": header,
                "sample": sample,
            })
        return result

    wb, stream = _openpyxl_load(path, True, True)
    try:
        names = [sheet] if sheet else wb.sheetnames
        for name in names:
            if name not in wb.sheetnames:
                raise ValueError(f"sheet not found: {name}")
            ws = wb[name]
            rows = ws.max_row or 0
            cols = ws.max_column or 0
            header: list = []
            sample: list = []
            if rows > 0 and cols > 0:
                preview = list(ws.iter_rows(min_row=1, max_row=min(rows, 6), values_only=True))
                if preview:
                    header = [_clean_value(v) for v in preview[0]]
                    sample = _clean_matrix(preview[1:6])
            result["sheets"].append({
                "name": name,
                "rows": rows,
                "columns": cols,
                "header": header,
                "sample": sample,
            })
        return result
    finally:
        wb.close()
        if stream is not None:
            stream.close()


def read_range(path: str, sheet: str | None = None, cell_range: str | None = None,
               offset: int = 0, limit: int = 100, show_formula: bool = False) -> dict:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    if limit <= 0:
        limit = 1
    _check_format(path)
    kind = _detect_kind(path)
    min_row, min_col, max_row, max_col = 1, 1, None, None
    if cell_range:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)

    if kind == "csv":
        if sheet and sheet != os.path.splitext(os.path.basename(path))[0]:
            raise ValueError("csv files have a single sheet")
        rows = _read_csv_rows(path)
        total_rows = len(rows)
        width = max((len(r) for r in rows), default=1)
        max_row = max_row or total_rows
        max_col = max_col or width
        start = min_row + max(offset, 0)
        end = min(max_row, start + limit - 1)
        values: list[list] = []
        for r in range(max(1, start), end + 1):
            row = rows[r - 1] if 0 <= r - 1 < total_rows else []
            values.append([_clean_value(row[c - 1]) if c - 1 < len(row) else None
                           for c in range(min_col, max_col + 1)])
        return {
            "path": path,
            "sheet": os.path.splitext(os.path.basename(path))[0],
            "range": cell_range or f"A1:{get_column_letter(max_col)}{max_row}",
            "offset": offset,
            "returned": len(values),
            "total_rows": total_rows,
            "values": values,
        }

    if kind == "xls":
        book = xlrd.open_workbook(path)
        ws = _xls_sheet(book, sheet)
        max_row = max_row or ws.nrows
        max_col = max_col or ws.ncols
        start = min_row + max(offset, 0)
        end = min(max_row, start + limit - 1)
        values = []
        for r in range(max(0, start - 1), min(ws.nrows, end)):
            values.append([_clean_value(_xls_cell_value(book, ws.cell(r, c)))
                           for c in range(max(0, min_col - 1), min(ws.ncols, max_col))])
        return {
            "path": path,
            "sheet": ws.name,
            "range": cell_range or f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
            "offset": offset,
            "returned": len(values),
            "total_rows": (max_row - min_row + 1),
            "values": values,
        }

    wb, stream = _openpyxl_load(path, True, not show_formula)
    try:
        ws = wb[sheet] if sheet else wb.active
        max_row = max_row or (ws.max_row or 1)
        max_col = max_col or (ws.max_column or 1)
        start = min_row + max(offset, 0)
        end = min(max_row, start + limit - 1)
        values = []
        if start <= max_row:
            for row in ws.iter_rows(min_row=start, max_row=end, min_col=min_col,
                                    max_col=max_col, values_only=True):
                values.append([_clean_value(v) for v in row])
        return {
            "path": path,
            "sheet": ws.title,
            "range": cell_range or f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}",
            "offset": offset,
            "returned": len(values),
            "total_rows": max_row - min_row + 1,
            "values": values,
        }
    finally:
        wb.close()
        if stream is not None:
            stream.close()


def _unique_columns(raw_header) -> list[str]:
    seen: dict[str, int] = {}
    columns: list[str] = []
    for index, value in enumerate(raw_header):
        name = str(value).strip() if value is not None else ""
        if not name:
            name = f"col{index + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        columns.append(name)
    return columns


def _rows_for_frame(path: str, sheet: str | None) -> list[list]:
    kind = _detect_kind(path)
    if kind == "csv":
        return _read_csv_rows(path)
    if kind == "xls":
        book = xlrd.open_workbook(path)
        ws = _xls_sheet(book, sheet)
        return [[_xls_cell_value(book, ws.cell(r, c)) for c in range(ws.ncols)]
                for r in range(ws.nrows)]
    wb, stream = _openpyxl_load(path, True, True)
    try:
        ws = wb[sheet] if sheet else wb.active
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
        if stream is not None:
            stream.close()


def load_sheet_frame(path: str, sheet: str | None = None) -> pd.DataFrame:
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    rows = _rows_for_frame(path, sheet)
    if not rows:
        return pd.DataFrame()
    columns = _unique_columns(rows[0])
    width = len(columns)
    data = []
    for row in rows[1:]:
        values = list(row)
        if len(values) < width:
            values += [None] * (width - len(values))
        data.append(values[:width])
    return pd.DataFrame(data, columns=columns)


def query(path: str, sql: str, sheet: str | None = None, limit: int = 1000) -> dict:
    frame = load_sheet_frame(path, sheet)
    connection = duckdb.connect()
    try:
        connection.execute("SET enable_external_access = false")
        connection.register("data", frame)
        cursor = connection.execute(sql)
        columns = [description[0] for description in cursor.description] if cursor.description else []
        rows = cursor.fetchmany(max(limit, 1))
        return {
            "columns": columns,
            "rows": [[_clean_value(v) for v in row] for row in rows],
            "row_count": len(rows),
            "limited_to": limit,
        }
    finally:
        connection.close()


def write_range(path: str, sheet: str, start_cell: str, values: list[list]) -> dict:
    if not values:
        raise ValueError("values must not be empty")
    _check_format(path)
    ext = _ext(path)
    if os.path.exists(path):
        kind = _detect_kind(path)
    elif ext in XLRD_FORMATS:
        kind = "xls"
    elif ext in CSV_FORMATS:
        kind = "csv"
    else:
        kind = "xlsx"
    row0, col0 = coordinate_to_tuple(start_cell)

    if kind == "xls":
        raise ValueError("writing .xls (BIFF) is not supported - save as xlsx or use the WPS live bridge")

    if kind == "csv":
        rows = _read_csv_rows(path) if os.path.exists(path) else []
        needed_rows = row0 - 1 + len(values)
        while len(rows) < needed_rows:
            rows.append([])
        written = 0
        for r_index, row in enumerate(values):
            target = rows[row0 - 1 + r_index]
            needed_cols = col0 - 1 + len(row)
            while len(target) < needed_cols:
                target.append("")
            for c_index, value in enumerate(row):
                target[col0 - 1 + c_index] = "" if value is None else value
                written += 1
        _atomic_csv_write(path, rows)
        return {"saved": path, "sheet": os.path.splitext(os.path.basename(path))[0],
                "start_cell": start_cell, "written": written}
    keep_vba = ext in {".xlsm", ".xltm"}
    if os.path.exists(path):
        wb, stream = _openpyxl_load(path, False, False, keep_vba=keep_vba)
        if stream is not None:
            stream.close()
    else:
        wb = openpyxl.Workbook()
        default = wb.active
        if default is not None and default.title != sheet:
            default.title = sheet
    try:
        if sheet not in wb.sheetnames:
            wb.create_sheet(sheet)
        ws = wb[sheet]
        written = 0
        for r_index, row in enumerate(values):
            for c_index, value in enumerate(row):
                cell = ws.cell(row=row0 + r_index, column=col0 + c_index)
                cell.value = value
                written += 1
        _save_edit(wb, path)
        return {"saved": path, "sheet": sheet, "start_cell": start_cell, "written": written}
    finally:
        wb.close()


def create_workbook(path: str, sheets: list[dict]) -> dict:
    if not sheets:
        raise ValueError("sheets must not be empty")
    ext = _check_format(path)

    if ext in XLRD_FORMATS:
        raise ValueError("creating .xls is not supported - use xlsx")

    if ext in CSV_FORMATS:
        if len(sheets) > 1:
            raise ValueError("csv supports a single sheet")
        values = sheets[0].get("values") or []
        _atomic_csv_write(path, [["" if v is None else v for v in row] for row in values])
        return {"created": path, "sheets": [os.path.splitext(os.path.basename(path))[0]]}

    wb = openpyxl.Workbook()
    try:
        default = wb.active
        created: list[str] = []
        for index, spec in enumerate(sheets):
            name = str(spec.get("name") or f"Sheet{index + 1}")
            values = spec.get("values") or []
            if index == 0 and default is not None:
                ws = default
                ws.title = name
            else:
                ws = wb.create_sheet(name)
            for r_index, row in enumerate(values):
                for c_index, value in enumerate(row):
                    ws.cell(row=r_index + 1, column=c_index + 1).value = value
            created.append(name)
        _save_edit(wb, path)
        return {"created": path, "sheets": created}
    finally:
        wb.close()


def convert(path: str, sheet: str | None = None, to: str = "parquet") -> dict:
    to = (to or "parquet").lower()
    if to not in ("parquet", "csv"):
        raise ValueError("to must be 'parquet' or 'csv'")
    frame = load_sheet_frame(path, sheet)
    base, _ = os.path.splitext(path)
    output = base + ("." + to)
    escaped = output.replace("'", "''")
    connection = duckdb.connect()
    try:
        connection.register("data", frame)
        if to == "parquet":
            connection.execute(f"COPY data TO '{escaped}' (FORMAT PARQUET)")
        else:
            connection.execute(f"COPY data TO '{escaped}' (FORMAT CSV, HEADER)")
    finally:
        connection.close()
    return {"output": output, "rows": len(frame), "columns": list(frame.columns), "format": to}


def _load_for_edit(path: str):
    if not os.path.exists(path):
        raise FileNotFoundError(path)
    _check_format(path)
    kind = _detect_kind(path)
    if kind == "xls":
        raise ValueError("editing .xls (BIFF) is not supported - save as xlsx or use the WPS live bridge")
    if kind == "csv":
        raise ValueError("this operation requires a workbook format (xlsx/xlsm), not csv/tsv")
    ext = _ext(path)
    keep_vba = ext in {".xlsm", ".xltm"}
    wb, stream = _openpyxl_load(path, False, False, keep_vba=keep_vba)
    if stream is not None:
        stream.close()
    return wb


def _remove_quietly(temp_path: str) -> None:
    try:
        os.remove(temp_path)
    except OSError:
        pass


def _replace_file(temp_path: str, path: str) -> None:
    try:
        os.replace(temp_path, path)
    except PermissionError as exc:
        _remove_quietly(temp_path)
        raise PermissionError(
            f"cannot write '{path}': file is locked (open in WPS/Excel?) - close it first or use the WPS live bridge"
        ) from exc
    except Exception:
        _remove_quietly(temp_path)
        raise


def _make_temp_path(path: str) -> str:
    directory = os.path.dirname(os.path.abspath(path))
    handle, temp_path = tempfile.mkstemp(
        prefix=".~" + os.path.basename(path) + ".", suffix=".tmp", dir=directory
    )
    os.close(handle)
    return temp_path


def _atomic_csv_write(path: str, rows) -> None:
    temp_path = _make_temp_path(path)
    try:
        with open(temp_path, "w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerows(rows)
    except Exception:
        _remove_quietly(temp_path)
        raise
    _replace_file(temp_path, path)


def _save_edit(wb, path: str) -> None:
    temp_path = _make_temp_path(path)
    try:
        wb.save(temp_path)
    except Exception:
        _remove_quietly(temp_path)
        raise
    _replace_file(temp_path, path)


def create_sheet(path: str, sheet: str, overwrite: bool = False) -> dict:
    wb = _load_for_edit(path)
    try:
        if sheet in wb.sheetnames:
            if not overwrite:
                raise ValueError(f"sheet already exists: {sheet}")
            del wb[sheet]
        wb.create_sheet(sheet)
        _save_edit(wb, path)
        return {"path": path, "sheet": sheet, "sheets": list(wb.sheetnames)}
    finally:
        wb.close()


def copy_sheet(path: str, src_sheet: str, dst_sheet: str) -> dict:
    wb = _load_for_edit(path)
    try:
        if src_sheet not in wb.sheetnames:
            raise ValueError(f"sheet not found: {src_sheet}")
        if dst_sheet in wb.sheetnames:
            raise ValueError(f"sheet already exists: {dst_sheet}")
        target = wb.copy_worksheet(wb[src_sheet])
        target.title = dst_sheet
        _save_edit(wb, path)
        return {"path": path, "src": src_sheet, "dst": dst_sheet, "sheets": list(wb.sheetnames)}
    finally:
        wb.close()


def rename_sheet(path: str, sheet: str, new_name: str) -> dict:
    wb = _load_for_edit(path)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"sheet not found: {sheet}")
        if new_name in wb.sheetnames:
            raise ValueError(f"sheet already exists: {new_name}")
        wb[sheet].title = new_name
        _save_edit(wb, path)
        return {"path": path, "old": sheet, "new": new_name, "sheets": list(wb.sheetnames)}
    finally:
        wb.close()


def delete_sheet(path: str, sheet: str) -> dict:
    wb = _load_for_edit(path)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"sheet not found: {sheet}")
        if len(wb.sheetnames) <= 1:
            raise ValueError("cannot delete the only sheet in the workbook")
        del wb[sheet]
        _save_edit(wb, path)
        return {"path": path, "deleted": sheet, "sheets": list(wb.sheetnames)}
    finally:
        wb.close()


_BORDER_ALIASES = {
    "thin": "thin", "continuous": "thin", "medium": "medium", "thick": "thick",
    "dash": "dashed", "dashed": "dashed", "dot": "dotted", "dotted": "dotted",
    "double": "double", "hair": "hair",
    "mediumdash": "mediumDashed", "mediumdashed": "mediumDashed",
    "dashdot": "dashDot", "mediumdashdot": "mediumDashDot",
    "dashdotdot": "dashDotDot", "mediumdashdotdot": "mediumDashDotDot",
    "slantdashdot": "slantDashDot", "none": None,
}

_UNDERLINE_ALIASES = {
    "none": None, "single": "single", "double": "double",
    "singleaccounting": "singleAccounting", "doubleaccounting": "doubleAccounting",
}

_VERT_ALIGN_ALIASES = {
    "baseline": None, "none": None, "superscript": "superscript", "subscript": "subscript",
}


def _normalize_rgb(value) -> str:
    if isinstance(value, (list, tuple)):
        value = value[0] if value else None
    if not value:
        raise ValueError("invalid color value (expected '#RRGGBB')")
    text = str(value).strip().lstrip("#").upper()
    if len(text) == 6:
        return text
    if len(text) == 8:
        return text[2:]
    raise ValueError(f"invalid color '{value}' (expected #RRGGBB)")


def _apply_cell_style(cell, style: dict) -> None:
    font_spec = style.get("font")
    if isinstance(font_spec, dict):
        font = copy(cell.font)
        if "bold" in font_spec:
            font.bold = bool(font_spec["bold"])
        if "italic" in font_spec:
            font.italic = bool(font_spec["italic"])
        if "strike" in font_spec:
            font.strike = bool(font_spec["strike"])
        if font_spec.get("size") is not None:
            font.size = float(font_spec["size"])
        if font_spec.get("name"):
            font.name = str(font_spec["name"])
        if font_spec.get("color"):
            font.color = _normalize_rgb(font_spec["color"])
        if "underline" in font_spec:
            raw = font_spec["underline"]
            if isinstance(raw, bool):
                font.underline = "single" if raw else None
            else:
                font.underline = _UNDERLINE_ALIASES.get(str(raw).lower(), str(raw))
        if "vertAlign" in font_spec:
            font.vertAlign = _VERT_ALIGN_ALIASES.get(str(font_spec["vertAlign"]).lower())
        cell.font = font

    fill_spec = style.get("fill")
    if isinstance(fill_spec, dict):
        pattern = str(fill_spec.get("pattern") or "solid")
        color = fill_spec.get("color") or fill_spec.get("fgColor")
        cell.fill = PatternFill(patternType=pattern, fgColor=_normalize_rgb(color))

    border_spec = style.get("border")
    if isinstance(border_spec, list):
        border = copy(cell.border)
        for item in border_spec:
            if not isinstance(item, dict):
                continue
            edge = str(item.get("type") or "").lower()
            side_style = _BORDER_ALIASES.get(str(item.get("style") or "thin").lower(), "thin")
            side = Side(style=side_style, color=_normalize_rgb(item.get("color") or "#000000"))
            if edge in ("left", "right", "top", "bottom"):
                setattr(border, edge, side)
            elif edge == "diagonal":
                border.diagonal = side
            elif edge == "diagonalup":
                border.diagonal = side
                border.diagonalUp = True
            elif edge == "diagonaldown":
                border.diagonal = side
                border.diagonalDown = True
        cell.border = border

    alignment_spec = style.get("alignment")
    if isinstance(alignment_spec, dict):
        alignment = copy(cell.alignment)
        if "horizontal" in alignment_spec:
            alignment.horizontal = alignment_spec["horizontal"]
        if "vertical" in alignment_spec:
            alignment.vertical = alignment_spec["vertical"]
        if "wrap" in alignment_spec:
            alignment.wrap_text = bool(alignment_spec["wrap"])
        if alignment_spec.get("rotation") is not None:
            alignment.textRotation = int(alignment_spec["rotation"])
        cell.alignment = alignment

    num_fmt = style.get("numFmt")
    decimal_places = style.get("decimalPlaces")
    if num_fmt:
        cell.number_format = str(num_fmt)
    elif decimal_places is not None:
        places = max(0, min(30, int(decimal_places)))
        cell.number_format = "0" if places == 0 else "0." + ("0" * places)


def format_range(path: str, sheet: str, cell_range: str, styles) -> dict:
    wb = _load_for_edit(path)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"sheet not found: {sheet}")
        ws = wb[sheet]
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        row_count = max_row - min_row + 1
        col_count = max_col - min_col + 1

        single: dict | None = None
        matrix: list | None = None
        if isinstance(styles, dict):
            single = styles
        elif isinstance(styles, list):
            if len(styles) != row_count or any(
                (not isinstance(row, list)) or len(row) != col_count for row in styles
            ):
                raise ValueError("styles matrix size must match the range size")
            matrix = styles
        else:
            raise ValueError("styles must be an object or a 2D array")

        styled = 0
        for r in range(row_count):
            for c in range(col_count):
                spec = single if single is not None else matrix[r][c]
                if not isinstance(spec, dict):
                    continue
                _apply_cell_style(ws.cell(row=min_row + r, column=min_col + c), spec)
                styled += 1
        _save_edit(wb, path)
        return {"path": path, "sheet": sheet, "range": cell_range, "styled_cells": styled}
    finally:
        wb.close()


def create_table(path: str, sheet: str, cell_range: str, table_name: str) -> dict:
    if not re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", table_name or ""):
        raise ValueError("table_name must start with a letter/underscore and contain only letters, digits, underscores")
    wb = _load_for_edit(path)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(f"sheet not found: {sheet}")
        ws = wb[sheet]
        existing = list(ws.tables.keys()) if hasattr(ws, "tables") else []
        if table_name in existing:
            raise ValueError(f"table already exists: {table_name}")
        table = Table(displayName=table_name, ref=cell_range)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        ws.add_table(table)
        _save_edit(wb, path)
        return {"path": path, "sheet": sheet, "range": cell_range, "table": table_name,
                "tables": list(ws.tables.keys())}
    finally:
        wb.close()
