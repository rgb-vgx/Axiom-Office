import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from excel_mcp import file_tools


class FileToolsTest(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix="exceltest_")
        self.xlsx = os.path.join(self.dir, "book.xlsx")

    def tearDown(self):
        for name in os.listdir(self.dir):
            try:
                os.remove(os.path.join(self.dir, name))
            except OSError:
                pass
        try:
            os.rmdir(self.dir)
        except OSError:
            pass

    def _create(self):
        file_tools.create_workbook(self.xlsx, [
            {"name": "Data", "values": [["Ten", "Diem"], ["An", 9.5], ["Binh", 8]]}
        ])

    def test_create_profile_read(self):
        self._create()
        prof = file_tools.profile(self.xlsx)
        self.assertEqual([s["name"] for s in prof["sheets"]], ["Data"])
        self.assertEqual(prof["sheets"][0]["header"], ["Ten", "Diem"])
        values = file_tools.read_range(self.xlsx, "Data", "A1:B3")["values"]
        self.assertEqual(values[1], ["An", 9.5])
        self.assertEqual(values[2], ["Binh", 8])

    def test_show_formula(self):
        self._create()
        file_tools.write_range(self.xlsx, "Data", "C1", [["=B2*2"]])
        formulas = file_tools.read_range(self.xlsx, "Data", "C1", show_formula=True)["values"]
        self.assertEqual(formulas, [["=B2*2"]])
        cached = file_tools.read_range(self.xlsx, "Data", "C1")["values"]
        self.assertEqual(cached, [[None]])

    def test_ragged_write_pads_with_null(self):
        self._create()
        file_tools.write_range(self.xlsx, "Data", "E1", [["a"], ["b", "c", "d"]])
        values = file_tools.read_range(self.xlsx, "Data", "E1:G2")["values"]
        self.assertEqual(values, [["a", None, None], ["b", "c", "d"]])

    def test_format_and_table(self):
        self._create()
        result = file_tools.format_range(self.xlsx, "Data", "A1:B1", {
            "font": {"bold": True},
            "fill": {"pattern": "solid", "color": "#FFFF00"},
            "numFmt": "@",
        })
        self.assertEqual(result["styled_cells"], 2)
        result = file_tools.create_table(self.xlsx, "Data", "A1:B3", "BangDiem")
        self.assertEqual(result["tables"], ["BangDiem"])

        import openpyxl
        wb = openpyxl.load_workbook(self.xlsx)
        try:
            self.assertTrue(wb["Data"]["A1"].font.bold)
            self.assertEqual(wb["Data"]["A1"].fill.fgColor.rgb, "00FFFF00")
            self.assertIn("BangDiem", wb["Data"].tables)
        finally:
            wb.close()

    def test_sheet_ops(self):
        self._create()
        file_tools.create_sheet(self.xlsx, "Temp")
        file_tools.copy_sheet(self.xlsx, "Data", "DataCopy")
        file_tools.rename_sheet(self.xlsx, "Temp", "Temp2")
        result = file_tools.delete_sheet(self.xlsx, "Temp2")
        self.assertEqual(result["sheets"], ["Data", "DataCopy"])
        with self.assertRaises(ValueError):
            file_tools.rename_sheet(self.xlsx, "Data", "DataCopy")
        file_tools.delete_sheet(self.xlsx, "DataCopy")
        with self.assertRaises(ValueError):
            file_tools.delete_sheet(self.xlsx, "Data")

    def test_convert_and_query(self):
        self._create()
        conv = file_tools.convert(self.xlsx, "Data", "parquet")
        self.assertTrue(os.path.exists(conv["output"]))
        self.assertEqual(conv["rows"], 2)
        result = file_tools.query(self.xlsx, "SELECT COUNT(*) AS n FROM data")
        self.assertEqual(result["rows"], [[2]])
        result = file_tools.query(self.xlsx, "SELECT SUM(Diem) AS s FROM data")
        self.assertEqual(result["rows"], [[17.5]])

    def test_query_sandbox_blocks_filesystem(self):
        self._create()
        with self.assertRaises(Exception):
            file_tools.query(self.xlsx, "SELECT * FROM read_csv_auto('C:/Windows/win.ini') LIMIT 1")
        with self.assertRaises(Exception):
            file_tools.query(self.xlsx, "COPY data TO 'C:/Users/Public/evil.csv' (FORMAT CSV)")

    def test_locked_file_raises_and_leaves_no_temp(self):
        self._create()
        handle = open(self.xlsx, "r")
        try:
            with self.assertRaises(PermissionError):
                file_tools.write_range(self.xlsx, "Data", "D1", [["x"]])
        finally:
            handle.close()
        leftovers = [n for n in os.listdir(self.dir) if n.startswith(".~")]
        self.assertEqual(leftovers, [])

    def test_csv_roundtrip(self):
        csv_path = os.path.join(self.dir, "data.csv")
        file_tools.create_workbook(csv_path, [{"name": "S", "values": [["A", "B"], [1, 2]]}])
        file_tools.write_range(csv_path, "S", "C1", [["C"]])
        values = file_tools.read_range(csv_path, None, "A1:C2")["values"]
        self.assertEqual(values, [["A", "B", "C"], [1, 2, None]])


if __name__ == "__main__":
    unittest.main()
