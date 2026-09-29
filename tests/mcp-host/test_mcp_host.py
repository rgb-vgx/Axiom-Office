"""Kiểm tra MCP server C# (AxiomOffice.Host.exe mcp) qua MCP client chính thức.

- Giao thức: initialize / tools/list / tools/call.
- Parity: so kết quả tool file của bản C# với tools/*-mcp/*/file_tools.py (bản Python cũ)
  trên cùng file (file do C# tạo, do python-docx/openpyxl/python-pptx tạo, file Office thật).
- Hợp lệ: đọc lại file C# ghi ra bằng python-docx / openpyxl / python-pptx.

Chạy (cần 3 venv của tools/*-mcp đã cài requirements):
    tools\\excel-mcp\\.venv\\Scripts\\python.exe tests\\mcp-host\\test_mcp_host.py [thư_mục_output]
"""
from __future__ import annotations

import asyncio
import glob
import json
import os
import shutil
import sys
import tempfile
import zipfile

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
EXE = os.path.join(ROOT, "src", "AxiomOffice", "bin", "Release", "AxiomOffice.Host.exe")

# Thư viện Python của 3 venv + package file_tools cũ để so parity.
for name in ("word-mcp", "ppt-mcp", "excel-mcp"):
    sys.path.insert(0, os.path.join(ROOT, "tools", name))
    site = os.path.join(ROOT, "tools", name, ".venv", "Lib", "site-packages")
    if os.path.isdir(site) and site not in sys.path:
        sys.path.append(site)

import docx  # noqa: E402
import openpyxl  # noqa: E402
import pptx  # noqa: E402
from mcp import ClientSession, StdioServerParameters  # noqa: E402
from mcp.client.stdio import stdio_client  # noqa: E402

from excel_mcp import file_tools as py_excel  # noqa: E402
from ppt_mcp import file_tools as py_ppt  # noqa: E402
from word_mcp import file_tools as py_word  # noqa: E402

RESULTS: list[tuple[bool, str, str]] = []


def attr(obj, *names):
    """mcp 2.x dùng snake_case (server_info), bản cũ dùng camelCase (serverInfo)."""
    for name in names:
        if hasattr(obj, name):
            return getattr(obj, name)
    raise AttributeError(names[0])


def check(ok: bool, name: str, detail: str = "") -> None:
    RESULTS.append((bool(ok), name, detail))
    print(("PASS " if ok else "FAIL ") + name + ("" if ok or not detail else "\n     " + detail[:900]))


def normalize(value):
    """So sánh bỏ qua khác biệt int/float (9 vs 9.0) và tuple/list."""
    if isinstance(value, dict):
        return {k: normalize(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [normalize(v) for v in value]
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def same(name: str, cs_value, py_value) -> None:
    a, b = normalize(cs_value), normalize(json.loads(json.dumps(py_value, default=str, ensure_ascii=False)))
    check(a == b, name, "C#: " + json.dumps(a, ensure_ascii=False)[:420] + "\n     Py: " + json.dumps(b, ensure_ascii=False)[:420])


def extra_samples(out: str, ext: str) -> list[str]:
    """File do Office/WPS thật tạo (real_*.ext trong thư mục output hoặc thư mục cha: phase*_*.docx)."""
    found = glob.glob(os.path.join(out, "real_*." + ext))
    if ext == "docx":
        found += glob.glob(os.path.join(os.path.dirname(out), "phase*_*.docx"))
    return [p for p in sorted(found) if os.path.getsize(p) > 0 and not os.path.basename(p).startswith("~$")]


class Client:
    def __init__(self, session: ClientSession):
        self.session = session

    async def call(self, tool: str, **arguments):
        result = await self.session.call_tool(tool, arguments)
        text = result.content[0].text if result.content else ""
        data = json.loads(text) if text else None
        return data, bool(attr(result, "is_error", "isError"))

    async def ok(self, tool: str, **arguments):
        data, is_error = await self.call(tool, **arguments)
        check(not is_error, "call " + tool, json.dumps(data, ensure_ascii=False))
        return data


async def test_protocol(session: ClientSession) -> None:
    tools = (await session.list_tools()).tools
    names = {t.name for t in tools}
    check(len(tools) == 50, "tools/list returns 50 tools", str(len(tools)))
    expected_py = set()
    for server in ("word-mcp/word_mcp", "ppt-mcp/ppt_mcp", "excel-mcp/excel_mcp"):
        with open(os.path.join(ROOT, "tools", server.split("/")[0], server.split("/")[1], "server.py"), encoding="utf-8") as handle:
            for line in handle:
                if line.startswith("def ") and "_safe" not in line and "_clean" not in line and "_live" not in line and "main" not in line:
                    expected_py.add(line[4:line.index("(")])
    missing = expected_py - names - {"excel_query"}
    check(not missing, "every Python tool exists in C# (except excel_query)", str(sorted(missing)))
    schema = attr(next(t for t in tools if t.name == "excel_format_range"), "input_schema", "inputSchema")
    check(schema.get("required") == ["path", "sheet", "cell_range", "styles"], "excel_format_range schema required", json.dumps(schema))


async def test_word(c: Client, out: str) -> None:
    path = os.path.join(out, "cs_word.docx")
    await c.ok("doc_create", path=path, overwrite=True, title="Báo cáo", author="Tester", paragraphs=[
        "Xin chào thế giới",
        {"text": "Tiêu đề chính", "style": "Heading 1"},
        {"text": "Dòng 1\nDòng 2\tcó tab", "style": "List Bullet"},
        {"table": [["Họ tên", "Điểm"], ["An", 9.5], ["Bình", None]]},
        {"text": "Kết thúc"},
    ])
    data, is_error = await c.call("doc_create", path=path, paragraphs=["x"])
    check(is_error and "exists" in data["error"], "doc_create refuses to overwrite without flag", json.dumps(data))
    data, is_error = await c.call("doc_create", path=os.path.join(out, "bad_style.docx"), overwrite=True, paragraphs=[{"text": "x", "style": "No Such Style"}])
    check(is_error, "doc_create unknown style -> error", json.dumps(data))

    d = docx.Document(path)
    check([p.text for p in d.paragraphs] == ["Xin chào thế giới", "Tiêu đề chính", "Dòng 1\nDòng 2\tcó tab", "Kết thúc"], "python-docx reads C# paragraphs",
          str([p.text for p in d.paragraphs]))
    check([p.style.name for p in d.paragraphs][:3] == ["Normal", "Heading 1", "List Bullet"], "python-docx sees styles", str([p.style.name for p in d.paragraphs]))
    check(d.tables[0].cell(1, 1).text == "9.5" and d.tables[0].style.name == "Table Grid", "python-docx reads C# table")
    check(d.core_properties.title == "Báo cáo" and d.core_properties.author == "Tester", "core properties written")

    samples = [path]
    py_path = os.path.join(out, "py_word.docx")
    doc = docx.Document()
    doc.add_heading("Tài liệu Python", 1)
    doc.add_paragraph("Đoạn có ").add_run("chữ đậm").bold = True
    table = doc.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Gộp ngang"
    table.cell(1, 2).merge(table.cell(2, 2))
    table.cell(1, 2).text = "Gộp dọc"
    doc.add_section()
    doc.add_paragraph("Sau section", style="Quote")
    doc.save(py_path)
    samples.append(py_path)
    samples += extra_samples(out, "docx")

    for sample in samples:
        label = os.path.basename(sample)
        cs = await c.ok("doc_profile", path=sample)
        py = py_word.profile(sample)
        for key in ("paragraphs", "tables", "sections", "style_counts"):
            same("doc_profile[%s].%s" % (label, key), cs[key], py[key])
        same("doc_profile[%s].core.title/author" % label, [cs["core"]["title"], cs["core"]["author"]], [py["core"]["title"] or None, py["core"]["author"] or None])
        same("doc_get_text[%s]" % label, await c.ok("doc_get_text", path=sample, max_chars=500), py_word.get_text(sample, 500))
        same("doc_get_text no tables[%s]" % label, await c.ok("doc_get_text", path=sample, include_tables=False), py_word.get_text(sample, 0, False))
        same("doc_find_text[%s]" % label, await c.ok("doc_find_text", path=sample, query="t"), py_word.find_text(sample, "t"))
        if py["tables"]:
            same("doc_extract_table[%s]" % label, await c.ok("doc_extract_table", path=sample, index=0), py_word.extract_table(sample, 0))


async def test_ppt(c: Client, out: str) -> None:
    path = os.path.join(out, "cs_deck.pptx")
    await c.ok("ppt_create", path=path, overwrite=True, slides=[
        {"title": "Giới thiệu công ty", "bullets": ["Thành lập 2010", "200 nhân sự"]},
        "Chỉ có tiêu đề",
        {"title": "Trống", "layout": 6},
        {"title": "Tiêu đề slide", "layout": 0, "text": "Phụ đề"},
    ])
    added = await c.ok("ppt_add_slide_file", path=path, title="Kế hoạch Q4", bullets=["Mục 1", "Mục 2"], layout=1)
    check(added["slide_count"] == 5, "ppt_add_slide_file slide_count", json.dumps(added))
    prs = pptx.Presentation(path)
    check(len(prs.slides) == 5, "python-pptx opens C# deck", str(len(prs.slides)))
    check(prs.slides[0].shapes.title.text == "Giới thiệu công ty", "title text readable by python-pptx")

    py_path = os.path.join(out, "py_deck.pptx")
    py_ppt.create(py_path, [{"title": "Python", "bullets": ["a", "b"]}, {"title": "Hai", "layout": 5}], overwrite=True)
    py_prs = pptx.Presentation(py_path)
    py_prs.slides[0].notes_slide.notes_text_frame.text = "Ghi chú thuyết trình"
    py_prs.save(py_path)
    for sample in [path, py_path] + extra_samples(out, "pptx"):
        label = os.path.basename(sample)
        same("ppt_profile[%s]" % label, await c.ok("ppt_profile", path=sample), py_ppt.profile(sample))
        same("ppt_get_text[%s]" % label, await c.ok("ppt_get_text", path=sample, max_chars=300), py_ppt.get_text(sample, 300))


async def test_excel(c: Client, out: str) -> None:
    import datetime
    from openpyxl.chart import BarChart, Reference

    path = os.path.join(out, "cs_book.xlsx")
    if os.path.exists(path):
        os.remove(path)
    await c.ok("excel_create", path=path, sheets=[
        {"name": "Data", "values": [["Tên", "Điểm", "Nhân đôi"], ["An", 9.5, "=B2*2"], ["Bình", 8, "=B3*2"], ["  có khoảng trắng ", True, None]]},
        {"name": "Khác", "values": [[1, 2, 3]]},
    ])
    wb = openpyxl.load_workbook(path)
    check(wb.sheetnames == ["Data", "Khác"], "openpyxl opens C# workbook", str(wb.sheetnames))
    check(wb["Data"]["C2"].value == "=B2*2" and wb["Data"]["A4"].value == "  có khoảng trắng " and wb["Data"]["B4"].value is True,
          "openpyxl reads values/formulas", str([wb["Data"]["C2"].value, wb["Data"]["A4"].value, wb["Data"]["B4"].value]))

    await c.ok("excel_write", path=path, sheet="Data", start_cell="D1", values=[["Ghi chú"], ["ok"]])
    await c.ok("excel_write", path=path, sheet="Mới", start_cell="B2", values=[["x", 1.25]])
    wb = openpyxl.load_workbook(path)
    check(wb["Data"]["D2"].value == "ok" and wb["Mới"]["C2"].value == 1.25, "excel_write updates + creates sheet")

    await c.ok("excel_format_range", path=path, sheet="Data", cell_range="A1:D1", styles={
        "font": {"bold": True, "color": "#FF0000", "size": 13, "name": "Arial", "underline": True},
        "fill": {"pattern": "solid", "color": "#FFFF00"},
        "border": [{"type": "bottom", "style": "thick", "color": "#0000FF"}],
        "alignment": {"horizontal": "center", "wrap": True},
    })
    await c.ok("excel_format_range", path=path, sheet="Data", cell_range="B2:B3", styles=[[{"decimalPlaces": 2}], [{"numFmt": "0.0%"}]])
    wb = openpyxl.load_workbook(path)
    a1 = wb["Data"]["A1"]
    check(a1.font.bold and a1.font.color.rgb.endswith("FF0000") and a1.font.sz == 13 and a1.font.name == "Arial" and a1.font.u == "single",
          "format font", str(a1.font))
    check(a1.fill.fgColor.rgb.endswith("FFFF00") and a1.border.bottom.style == "thick" and a1.alignment.horizontal == "center" and a1.alignment.wrap_text,
          "format fill/border/alignment")
    check(wb["Data"]["B2"].number_format == "0.00" and wb["Data"]["B3"].number_format == "0.0%", "format numFmt",
          wb["Data"]["B2"].number_format + " / " + wb["Data"]["B3"].number_format)

    table = await c.ok("excel_create_table", path=path, sheet="Data", cell_range="A1:D4", table_name="BangDiem")
    check(table["tables"] == ["BangDiem"], "excel_create_table result", json.dumps(table))
    data, is_error = await c.call("excel_create_table", path=path, sheet="Data", cell_range="A1:D4", table_name="BangDiem")
    check(is_error, "duplicate table name rejected")
    wb = openpyxl.load_workbook(path)
    check("BangDiem" in wb["Data"].tables, "openpyxl sees table")

    await c.ok("excel_copy_sheet", path=path, src_sheet="Data", dst_sheet="Data copy")
    await c.ok("excel_rename_sheet", path=path, sheet="Khác", new_name="Khác 2")
    await c.ok("excel_create_sheet", path=path, sheet="Trống")
    deleted = await c.ok("excel_delete_sheet", path=path, sheet="Mới")
    check(deleted["sheets"] == ["Data", "Khác 2", "Data copy", "Trống"], "sheet operations", json.dumps(deleted))
    wb = openpyxl.load_workbook(path)
    check(wb.sheetnames == ["Data", "Khác 2", "Data copy", "Trống"] and wb["Data copy"]["A1"].font.bold, "openpyxl sees sheet ops + copied style")
    data, is_error = await c.call("excel_delete_sheet", path=os.path.join(out, "one.xlsx"), sheet="x")
    check(is_error, "delete on missing file -> error")

    # File do openpyxl tạo: ngày tháng, công thức, chart -> đọc parity và chart phải còn sau khi C# ghi.
    py_path = os.path.join(out, "py_book.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Số liệu"
    ws.append(["Ngày", "Doanh thu", "Tỉ lệ", "Công thức"])
    ws.append([datetime.datetime(2024, 1, 15), 1200, 0.25, "=B2*C2"])
    ws.append([datetime.datetime(2024, 2, 1, 13, 30), 800.5, None, "=SUM(B2:B3)"])
    ws.append([datetime.time(8, 15), -3, "chữ", True])
    ws["B2"].number_format = "#,##0"
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    ws.add_chart(chart, "F2")
    wb.create_sheet("Phụ").append(["a", 1])
    wb.save(py_path)
    for sample in [py_path, path] + extra_samples(out, "xlsx"):
        label = os.path.basename(sample)
        same("excel_profile[%s]" % label, await c.ok("excel_profile", path=sample), py_excel.profile(sample))
        same("excel_read[%s]" % label, await c.ok("excel_read", path=sample, cell_range="A1:D4"), py_excel.read_range(sample, None, "A1:D4"))
        same("excel_read formulas[%s]" % label, await c.ok("excel_read", path=sample, show_formula=True), py_excel.read_range(sample, None, None, 0, 100, True))
        same("excel_read paging[%s]" % label, await c.ok("excel_read", path=sample, offset=1, limit=2), py_excel.read_range(sample, None, None, 1, 2))
    await c.ok("excel_write", path=py_path, sheet="Số liệu", start_cell="E1", values=[["Ghi chú"]])
    with zipfile.ZipFile(py_path) as z:
        check(any(n.startswith("xl/charts/") for n in z.namelist()), "chart preserved after excel_write", str(z.namelist()))
    check(openpyxl.load_workbook(py_path)["Số liệu"]["E1"].value == "Ghi chú", "openpyxl reads write on openpyxl file")

    csv_path = os.path.join(out, "cs.csv")
    await c.ok("excel_create", path=csv_path, sheets=[{"values": [["Tên", "Điểm"], ["An", 9.5], ["Bình, Hà", 8]]}])
    await c.ok("excel_write", path=csv_path, sheet="cs", start_cell="C1", values=[["Ghi chú"], ["có \"ngoặc\""]])
    same("excel_profile[csv]", await c.ok("excel_profile", path=csv_path), py_excel.profile(csv_path))
    same("excel_read[csv]", await c.ok("excel_read", path=csv_path), py_excel.read_range(csv_path))
    converted = await c.ok("excel_convert", path=path, sheet="Data")
    check(os.path.exists(converted["output"]) and converted["columns"][:2] == ["Tên", "Điểm"], "excel_convert csv", json.dumps(converted))
    data, is_error = await c.call("excel_convert", path=path, to="parquet")
    check(is_error and "parquet" in data["error"], "parquet reports clear error")


async def main() -> int:
    out = sys.argv[1] if len(sys.argv) > 1 else tempfile.mkdtemp(prefix="mcp_host_")
    os.makedirs(out, exist_ok=True)
    print("exe:", EXE)
    print("out:", out)
    params = StdioServerParameters(command=EXE, args=["mcp"])
    async with stdio_client(params) as (read, write):
        async with ClientSession(read, write) as session:
            init = await session.initialize()
            info = attr(init, "server_info", "serverInfo")
            check(info.name == "office-tools", "initialize serverInfo", str(info))
            await test_protocol(session)
            client = Client(session)
            await test_word(client, out)
            await test_ppt(client, out)
            await test_excel(client, out)
    failed = [r for r in RESULTS if not r[0]]
    print("\n%d passed, %d failed" % (len(RESULTS) - len(failed), len(failed)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
